import discord
from openpyxl import Workbook, load_workbook
import os

intents = discord.Intents.default() # bot ko permissions deta hai ki wo messages ko dekh sake
intents.message_content = True # bot ko message ka content (jo user bhej raha hai) dikhai dega ..jiase $expense ya $view_expenses
bot = discord.Client(intents=intents) #Bot ko initialize karta hain, aur intents ko pass karte hain

@bot.event
async def on_ready():
    print(f'Logged in as {bot.user}')

@bot.event
async def on_message(message):
    if message.author == bot.user:
        return

    if message.content.startswith('$expense'):
        # Extract the expense details
        try:
            # Example: $expense 50 Coffee 2025-05-01
            parts = message.content.split(' ')
            amount = parts[1]
            place = parts[2]
            date = parts[3]

            # Get the user's ID to create a unique file
            user_id = str(message.author.id)

            # Path for storing user's Excel file
            file_path = f'{user_id}_expenses.xlsx'

            # Check if the file already exists
            if os.path.exists(file_path):
                # Load the existing workbook
                wb = load_workbook(file_path)
                sheet = wb.active
            else:
                # Create a new workbook
                wb = Workbook()
                sheet = wb.active
                sheet.append(['Amount', 'Place', 'Date'])

            # Add the new expense
            sheet.append([amount, place, date])

            # Save the workbook
            wb.save(file_path)

            await message.channel.send(f'Expense added: {amount} at {place} on {date}')

        except Exception as e:
            await message.channel.send('Error: Invalid expense format! Please use the format "$expense <amount> <place> <date>"')

    if message.content.startswith('$view_expenses'):
        # Allow user to view only their own file
        try:
            user_id = str(message.author.id)
            file_path = f'{user_id}_expenses.xlsx'

            if os.path.exists(file_path):
                # Send the file to the user
                await message.channel.send(
                    content="Here is your expense file:", 
                    file=discord.File(file_path)
                )

            else:
                await message.channel.send("You have no expenses recorded yet.")
        
        except Exception as e:
            await message.channel.send('Error fetching expenses.')

bot.run('MTM2NzE4ODg1NjkxODUwNzYwMQ.GzcSji.uJ3QW9u1uf4WNC_STX3B8B-lfQdg4KYjWxVPm0')

#COMMANDS YE HAI :--- py bot.py
#$expense <amount> <place> <date>
#$view_expenses

